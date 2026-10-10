from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from docx import Document
from io import BytesIO
from pathlib import Path
import pickle
import fitz
import re

BASE_DIR = Path(__file__).resolve().parent
MODEL_FILE = BASE_DIR / "trained_model.pkl"

app = FastAPI(
    title="Placement Resume Analysis Service",
    version="0.3.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:5000"
    ],
    allow_methods=["*"],
    allow_headers=["*"]
)

try:
    with open(MODEL_FILE, "rb") as file:
        trained_model = pickle.load(file)
except Exception as error:
    raise RuntimeError(
        f"Could not load trained model: {error}"
    ) from error

SKILLS = {
    "java", "python", "javascript", "typescript", "react",
    "node.js", "express", "mongodb", "sql", "mysql",
    "postgresql", "spring boot", "rest api", "rest apis",
    "data structures", "algorithms", "machine learning",
    "deep learning", "nlp", "pandas", "numpy",
    "scikit-learn", "git", "github", "html", "css",
    "flutter", "dart", "firebase", "aws", "docker",
    "communication", "problem solving", "fastapi",
    "flask", "django", "c", "c++", "c#", ".net",
    "php", "kotlin", "swift", "react native",
    "angular", "vue", "redis", "postgres", "oracle",
    "linux", "windows", "api development", "oop",
    "object oriented programming", "data analysis",
    "data visualization", "excel", "power bi",
    "tableau", "pytorch", "tensorflow", "keras",
    "statistics", "scipy", "matplotlib", "seaborn",
    "jenkins", "kubernetes", "gitlab", "github actions",
    "agile", "scrum", "unit testing", "selenium",
    "cybersecurity", "networking", "ethical hacking",
    "cloud computing", "azure", "gcp"
}

SECTION_PATTERNS = {
    "contact": r"\b(contact|email|phone|mobile)\b",
    "education": r"\b(education|academic background|qualifications?)\b",
    "skills": r"\b(technical skills|skills|technical expertise|core competencies)\b",
    "projects": r"\b(projects?|personal projects|academic projects)\b",
    "experience": r"\b(experience|employment|internship|work history)\b",
    "certifications": r"\b(certifications?|certificates|licenses)\b",
    "summary": r"\b(professional summary|career objective|profile summary|objective)\b"
}


def extract_text(filename, content):
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    try:
        if ext == "pdf":
            with fitz.open(stream=content, filetype="pdf") as pdf:
                return "\n".join(page.get_text() for page in pdf)

        if ext == "docx":
            doc = Document(BytesIO(content))
            parts = [p.text for p in doc.paragraphs]

            for table in doc.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))

            return "\n".join(parts)

    except Exception:
        raise HTTPException(
            status_code=400,
            detail="Could not read the uploaded resume."
        )

    raise HTTPException(
        status_code=400,
        detail="Upload a PDF or DOCX resume."
    )


def find_skills(text):
    lower = text.lower()
    found = set()

    for skill in SKILLS:
        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(skill)
            + r"(?![a-z0-9])"
        )

        if re.search(pattern, lower):
            found.add(skill)

    if "rest apis" in found:
        found.discard("rest apis")
        found.add("rest api")

    if "postgres" in found:
        found.discard("postgres")
        found.add("postgresql")

    if "object oriented programming" in found:
        found.discard("object oriented programming")
        found.add("oop")

    return sorted(found)


def calculate_similarity(resume, job):
    if not resume.strip() or not job.strip():
        return 0.0

    try:
        vectorizer = TfidfVectorizer(stop_words="english")
        matrix = vectorizer.fit_transform([resume, job])
        score = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]

        return round(float(score) * 100, 2)

    except ValueError:
        return 0.0


def analyze_structure(text):
    checks = {
        section: bool(re.search(pattern, text, re.IGNORECASE))
        for section, pattern in SECTION_PATTERNS.items()
    }

    checks["email_detected"] = bool(
        re.search(
            r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
            text
        )
    )

    checks["phone_detected"] = bool(
        re.search(
            r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)",
            text
        )
    )

    return checks


def generate_suggestions(missing_skills, checks, text):
    suggestions = []

    if missing_skills:
        suggestions.append(
            "Review these job-related skills: "
            + ", ".join(missing_skills)
            + ". Include them only if you genuinely have the skills."
        )

    if not checks["education"]:
        suggestions.append("Consider adding a clearly labelled Education section.")

    if not checks["skills"]:
        suggestions.append("Add a Skills section listing your relevant technical abilities.")

    if not checks["projects"]:
        suggestions.append("Include relevant academic or personal projects with your contributions.")

    if not checks["experience"]:
        suggestions.append(
            "If applicable, include internships, work experience, or relevant practical experience."
        )

    if not checks["certifications"]:
        suggestions.append(
            "If you have relevant certifications, consider adding a Certifications section."
        )

    if not checks["summary"]:
        suggestions.append(
            "Consider adding a short professional summary tailored to the target role."
        )

    if not checks["email_detected"]:
        suggestions.append("Check that your resume contains a professional email address.")

    if not checks["phone_detected"]:
        suggestions.append("Check that your resume contains a readable contact phone number.")

    if len(text.split()) < 100:
        suggestions.append(
            "The extracted resume is relatively short. Check that all relevant sections were extracted correctly."
        )

    if not suggestions:
        suggestions.append(
            "Use specific examples and measurable achievements to support your relevant skills."
        )

    return suggestions


@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "resume-analysis",
        "model_loaded": trained_model is not None
    }


@app.post("/analyze")
async def analyze(
    resume: UploadFile = File(...),
    job_description: str = Form(...)
):
    filename = resume.filename or ""

    if not re.search(r"\.(pdf|docx)$", filename, re.IGNORECASE):
        raise HTTPException(
            status_code=400,
            detail="Upload a PDF or DOCX resume."
        )

    content = await resume.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="Uploaded file is empty."
        )

    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail="Resume must be 5 MB or smaller."
        )

    if not job_description.strip():
        raise HTTPException(
            status_code=400,
            detail="Job description cannot be empty."
        )

    text = extract_text(filename, content)

    if len(text.strip()) < 30:
        raise HTTPException(
            status_code=400,
            detail="Not enough text found. Scanned PDFs requiring OCR are not supported yet."
        )

    resume_skills = find_skills(text)
    job_skills = find_skills(job_description)

    matched_skills = sorted(set(resume_skills) & set(job_skills))
    missing_skills = sorted(set(job_skills) - set(resume_skills))

    skill_score = (
        round(len(matched_skills) / len(job_skills) * 100, 2)
        if job_skills else 0.0
    )

    text_score = calculate_similarity(text, job_description)
    checks = analyze_structure(text)

    structure_score = round(
        sum(checks.values()) / len(checks) * 100,
        2
    )

    heuristic_score = round(
    0.55 * skill_score
    + 0.30 * text_score
    + 0.15 * structure_score,
    2
)

    model_input = text + " [SEP] " + job_description

    prediction = int(trained_model.predict([model_input])[0])

    probabilities = trained_model.predict_proba([model_input])[0]
    classes = list(trained_model.classes_)
    relevant_index = classes.index(1)

    model_relevance_score = round(
        float(probabilities[relevant_index]) * 100,
        2
    )

    suggestions = generate_suggestions(
        missing_skills,
        checks,
        text
    )

    return {
        "ats_compatibility_score": heuristic_score,
        "score_type": "heuristic_not_official",
        "text_similarity_score": text_score,
        "skill_match_score": skill_score,
        "resume_structure_score": structure_score,
        "matched_skills": matched_skills,
        "missing_job_skills": missing_skills,
        "resume_skills": resume_skills,
        "job_skills_detected": job_skills,
        "structure_checks": checks,
        "suggestions": suggestions,
        "trained_model_prediction": (
            "relevant" if prediction == 1 else "less_relevant"
        ),
        "trained_model_relevance_score": model_relevance_score,
        "trained_model_note": (
            "Experimental model trained on synthetic examples. "
            "Its prediction and probability are not validated measures "
            "of candidate suitability."
        ),
        "note": (
            "The compatibility score is a heuristic based on text similarity, "
            "recognized skills, and section checks. It is not an official ATS "
            "score or a hiring recommendation."
        )
    }