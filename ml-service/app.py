from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from docx import Document
from io import BytesIO
from pathlib import Path
import pickle
import fitz
import re

BASE_DIR = Path(__file__).resolve().parent
MODEL_FILE = BASE_DIR / "trained_model.pkl"

app = FastAPI(title="Placement Resume Analysis Service", version="0.5.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:5000",
    ],
    allow_methods=["*"],
    allow_headers=["*"],
)

with open(MODEL_FILE, "rb") as file:
    trained_model = pickle.load(file)

SKILLS = {
    "java", "python", "javascript", "typescript", "react",
    "node.js", "express", "mongodb", "sql", "mysql",
    "postgresql", "spring boot", "rest api", "rest apis",
    "data structures", "algorithms", "machine learning",
    "deep learning", "nlp", "pandas", "numpy", "scikit-learn",
    "git", "github", "html", "css", "flutter", "dart", "firebase",
    "aws", "docker", "communication", "problem solving",
    "fastapi", "flask", "django", "c", "c++", "c#", ".net",
    "php", "kotlin", "swift", "react native", "angular", "vue",
    "redis", "postgres", "oracle", "linux", "windows",
    "api development", "oop", "object oriented programming",
    "data analysis", "data visualization", "excel", "power bi",
    "tableau", "pytorch", "tensorflow", "keras", "statistics",
    "scipy", "matplotlib", "seaborn", "jenkins", "kubernetes",
    "gitlab", "github actions", "agile", "scrum", "unit testing",
    "selenium", "cybersecurity", "networking", "ethical hacking",
    "cloud computing", "azure", "gcp", "hibernate", "maven",
    "pytest", "android", "state management", "responsive design",
    "network security", "siem", "incident response", "wireshark",
}

SECTION_PATTERNS = {
    "contact": r"\b(contact|email|phone|mobile)\b",
    "education": r"\b(education|academic background|qualifications?)\b",
    "skills": r"\b(technical skills|skills|technical expertise|core competencies)\b",
    "projects": r"\b(projects?|personal projects|academic projects)\b",
    "experience": r"\b(experience|employment|internship|work history)\b",
    "certifications": r"\b(certifications?|certificates|licenses)\b",
    "summary": r"\b(professional summary|career objective|profile summary|objective)\b",
}

REQUIRED_HEADING = re.compile(
    r"(?im)^\s*(?:required skills|mandatory skills|must[- ]have skills|must[- ]have)\s*:?\s*$"
)
PREFERRED_HEADING = re.compile(
    r"(?im)^\s*(?:preferred skills|good[- ]to[- ]have skills|good[- ]to[- ]have|optional skills)\s*:?\s*$"
)
ANY_SECTION_HEADING = re.compile(
    r"(?im)^\s*(?:required skills|mandatory skills|must[- ]have skills|must[- ]have|"
    r"preferred skills|good[- ]to[- ]have skills|good[- ]to[- ]have|optional skills|"
    r"responsibilities|job responsibilities|qualifications|requirements|about the role|"
    r"job description)\s*:?\s*$"
)


def extract_text(filename, content):
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    try:
        if ext == "pdf":
            with fitz.open(stream=content, filetype="pdf") as pdf:
                return "\n".join(page.get_text() for page in pdf)

        if ext == "docx":
            doc = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in doc.paragraphs]

            for table in doc.tables:
                for row in table.rows:
                    parts.append(" ".join(cell.text for cell in row.cells))

            return "\n".join(parts)

    except Exception as error:
        raise HTTPException(
            status_code=400,
            detail="Could not read the uploaded resume.",
        ) from error

    raise HTTPException(
        status_code=400,
        detail="Upload a PDF or DOCX resume.",
    )


def find_skills(text):
    lower = text.lower()
    found = set()

    for skill in SKILLS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.add(skill)

    if "rest apis" in found:
        found.discard("rest apis")
        found.add("rest api")

    if "postgres" in found:
        found.discard("postgres")
        found.add("postgresql")

    if "object oriented programming" in found:
        found.discard("object oriented programming")
        found.add("oop")

    return sorted(found)


def calculate_similarity(resume, job):
    if not resume.strip() or not job.strip():
        return 0.0

    try:
        vectorizer = TfidfVectorizer(stop_words="english")
        matrix = vectorizer.fit_transform([resume, job])
        score = cosine_similarity(matrix[0:1], matrix[1:2])[0][0]
        return round(float(score) * 100, 2)
    except ValueError:
        return 0.0


def analyze_structure(text):
    checks = {
        section: bool(re.search(pattern, text, re.IGNORECASE))
        for section, pattern in SECTION_PATTERNS.items()
    }
    checks["email_detected"] = bool(
        re.search(r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", text)
    )
    checks["phone_detected"] = bool(
        re.search(r"(?<!\w)(?:\+?\d[\d\s().-]{7,}\d)(?!\w)", text)
    )
    return checks


def extract_heading_skills(job_description, heading_pattern):
    matches = list(heading_pattern.finditer(job_description))
    skills = set()

    for match in matches:
        start = match.end()
        next_heading = ANY_SECTION_HEADING.search(job_description, start)
        end = next_heading.start() if next_heading else len(job_description)
        skills.update(find_skills(job_description[start:end]))

    return sorted(skills)


def get_skill_categories(job_description, all_job_skills):
    required = extract_heading_skills(job_description, REQUIRED_HEADING)
    preferred = extract_heading_skills(job_description, PREFERRED_HEADING)

    required_set = set(required)
    preferred_set = set(preferred) - required_set
    categorized = required_set | preferred_set

    # If headings are not provided, do not guess which skills are mandatory.
    uncategorized = set(all_job_skills) - categorized

    return sorted(required_set), sorted(preferred_set), sorted(uncategorized)


def generate_suggestions(missing_skills, checks, text):
    suggestions = []

    if missing_skills:
        suggestions.append(
            "Review these job-related skills: "
            + ", ".join(missing_skills)
            + ". Include them only if you genuinely have the skills."
        )

    if not checks["education"]:
        suggestions.append("Consider adding a clearly labelled Education section.")
    if not checks["skills"]:
        suggestions.append("Add a Skills section listing your relevant technical abilities.")
    if not checks["projects"]:
        suggestions.append("Include relevant academic or personal projects with your contributions.")
    if not checks["experience"]:
        suggestions.append("If applicable, include internships, work experience, or relevant practical experience.")
    if not checks["certifications"]:
        suggestions.append("If you have relevant certifications, consider adding a Certifications section.")
    if not checks["summary"]:
        suggestions.append("Consider adding a short professional summary tailored to the target role.")
    if not checks["email_detected"]:
        suggestions.append("Check that your resume contains a professional email address.")
    if not checks["phone_detected"]:
        suggestions.append("Check that your resume contains a readable contact phone number.")
    if len(text.split()) < 100:
        suggestions.append("The extracted resume is relatively short. Check that all relevant sections were extracted correctly.")
    if not suggestions:
        suggestions.append("Use specific examples and measurable achievements to support your relevant skills.")

    return suggestions


@app.get("/health")
def health():
    return {
        "status": "ok",
        "service": "resume-analysis",
        "model_loaded": trained_model is not None,
    }


@app.post("/analyze")
async def analyze(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
):
    filename = resume.filename or ""

    if not re.search(r"\.(pdf|docx)$", filename, re.IGNORECASE):
        raise HTTPException(status_code=400, detail="Upload a PDF or DOCX resume.")

    content = await resume.read()

    if not content:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Resume must be 5 MB or smaller.")
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description cannot be empty.")

    text = extract_text(filename, content)

    if len(text.strip()) < 30:
        raise HTTPException(
            status_code=400,
            detail="Not enough text found. Scanned PDFs requiring OCR are not supported yet.",
        )

    resume_skills = find_skills(text)
    job_skills = find_skills(job_description)
    resume_skill_set = set(resume_skills)
    job_skill_set = set(job_skills)

    matched_skills = sorted(resume_skill_set & job_skill_set)
    missing_skills = sorted(job_skill_set - resume_skill_set)

    skill_score = (
        round(len(matched_skills) / len(job_skills) * 100, 2)
        if job_skills else 0.0
    )

    required_skills, preferred_skills, uncategorized_skills = get_skill_categories(
        job_description, job_skills
    )

    matched_required = sorted(set(required_skills) & resume_skill_set)
    missing_required = sorted(set(required_skills) - resume_skill_set)
    matched_preferred = sorted(set(preferred_skills) & resume_skill_set)
    missing_preferred = sorted(set(preferred_skills) - resume_skill_set)
    matched_uncategorized = sorted(set(uncategorized_skills) & resume_skill_set)
    missing_uncategorized = sorted(set(uncategorized_skills) - resume_skill_set)

    required_score = (
        round(len(matched_required) / len(required_skills) * 100, 2)
        if required_skills else None
    )
    preferred_score = (
        round(len(matched_preferred) / len(preferred_skills) * 100, 2)
        if preferred_skills else None
    )
    uncategorized_score = (
        round(len(matched_uncategorized) / len(uncategorized_skills) * 100, 2)
        if uncategorized_skills else None
    )

    text_score = calculate_similarity(text, job_description)
    checks = analyze_structure(text)
    structure_score = round(sum(checks.values()) / len(checks) * 100, 2)

    if required_skills or preferred_skills:
        components = []
        if required_score is not None:
            components.append((0.70, required_score))
        if preferred_score is not None:
            components.append((0.20, preferred_score))
        if uncategorized_score is not None:
            components.append((0.10, uncategorized_score))

        weight_total = sum(weight for weight, _ in components)
        skill_category_score = sum(weight * score for weight, score in components) / weight_total
        compatibility_score = round(0.85 * skill_category_score + 0.15 * text_score, 2)
    elif job_skills:
        compatibility_score = round(0.80 * skill_score + 0.20 * text_score, 2)
    else:
        compatibility_score = text_score

    suggestions = generate_suggestions(missing_skills, checks, text)
    if missing_required:
        suggestions.append(
            "Required skills not detected: "
            + ", ".join(missing_required)
            + ". Include them only if you genuinely have those skills."
        )
    if missing_preferred:
        suggestions.append(
            "Preferred skills not detected: "
            + ", ".join(missing_preferred)
            + ". Consider these only if relevant to your experience."
        )

    model_input = text + " [SEP] " + job_description
    prediction = int(trained_model.predict([model_input])[0])
    probabilities = trained_model.predict_proba([model_input])[0]
    classes = list(trained_model.classes_)
    model_score = round(float(probabilities[classes.index(1)]) * 100, 2)

    return {
        "ats_compatibility_score": compatibility_score,
        "score_type": "heuristic_not_official",
        "text_similarity_score": text_score,
        "skill_match_score": skill_score,
        "resume_structure_score": structure_score,
        "matched_skills": matched_skills,
        "missing_job_skills": missing_skills,
        "resume_skills": resume_skills,
        "job_skills_detected": job_skills,
        "structure_checks": checks,
        "suggestions": suggestions,
        "required_skills": required_skills,
        "matched_required_skills": matched_required,
        "missing_required_skills": missing_required,
        "required_skill_score": required_score,
        "preferred_skills": preferred_skills,
        "matched_preferred_skills": matched_preferred,
        "missing_preferred_skills": missing_preferred,
        "preferred_skill_score": preferred_score,
        "uncategorized_job_skills": uncategorized_skills,
        "matched_uncategorized_skills": matched_uncategorized,
        "missing_uncategorized_skills": missing_uncategorized,
        "uncategorized_skill_score": uncategorized_score,
        "trained_model_prediction": "relevant" if prediction == 1 else "less_relevant",
        "trained_model_relevance_score": model_score,
        "trained_model_note": (
            "Experimental model trained on synthetic examples. "
            "Its prediction and probability are not validated measures "
            "of candidate suitability."
        ),
        
    }
